"""ONE-TIME migration of the existing customer registries from a CSV or XLSX file.

This is deliberately NOT part of the Admin Portal. Typical use:

    python scripts/import_registries.py customers.xlsx                 # dry run: validate only
    python scripts/import_registries.py customers.xlsx --commit        # write valid rows
    python scripts/import_registries.py customers.csv --commit --strict  # write nothing if ANY row is bad

Required columns (header names are case-insensitive; spaces/hyphens are treated as underscores):
    registry_number, registry_code, dedicated_to, dedicated_by, dedication_date
Optional columns:
    occasion, star_catalogue, constellation, distance, dedication_message, status (active|inactive)

Dates must match --date-format (default %Y-%m-%d) or be real Excel date cells. Ambiguous
day/month dates are never guessed. Certificates are not imported; upload them in the Admin Portal.
"""
import _bootstrap  # noqa: F401
import argparse
import csv
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from pydantic import ValidationError

from app.config import load_settings
from app.database.connection import connect, init_db
from app.database.registry_repository import DuplicateRegistryNumber, RegistryRepository
from app.schemas.registry import RegistryWrite
from app.services.crypto import CodeVault
from app.services.registry_service import to_storage

REQUIRED = ["registry_number", "registry_code", "dedicated_to", "dedicated_by", "dedication_date"]
OPTIONAL = ["occasion", "star_catalogue", "constellation", "distance", "dedication_message", "status"]
ALIASES = {
    "registry_no": "registry_number", "registry_#": "registry_number", "number": "registry_number",
    "code": "registry_code", "dedicated_for": "dedicated_to", "to": "dedicated_to", "from": "dedicated_by",
    "by": "dedicated_by", "date": "dedication_date", "date_of_dedication": "dedication_date",
    "message": "dedication_message", "catalogue": "star_catalogue", "star_catalog": "star_catalogue",
}


def normalise_header(h: object) -> str:
    key = str(h or "").strip().lower().replace(" ", "_").replace("-", "_")
    return ALIASES.get(key, key)


def read_rows(path: Path) -> tuple[list[str], list[tuple[int, dict]]]:
    """Return (headers, [(source_row_number, {header: value})])."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.reader(fh)
            raw = list(reader)
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        raw = [list(r) for r in wb.active.iter_rows(values_only=True)]
    else:
        raise SystemExit(f"Unsupported file type '{suffix}'. Use .csv or .xlsx.")
    if not raw:
        raise SystemExit("The file is empty.")
    headers = [normalise_header(h) for h in raw[0]]
    rows = []
    for i, values in enumerate(raw[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue  # skip blank lines
        rows.append((i, dict(zip(headers, values))))
    return headers, rows


def parse_date(value: object, fmt: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValueError("dedication_date is empty")
    try:
        return datetime.strptime(text, fmt).date()
    except ValueError:
        raise ValueError(f"dedication_date '{text}' does not match format {fmt}") from None


def to_text(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def build_record(row: dict, date_format: str) -> RegistryWrite:
    data = {k: to_text(row.get(k)) for k in REQUIRED + OPTIONAL if k != "dedication_date"}
    data["dedication_date"] = parse_date(row.get("dedication_date"), date_format)
    if not data.get("status"):
        data["status"] = "active"
    data["status"] = str(data["status"]).lower()
    return RegistryWrite(**{k: v for k, v in data.items() if v is not None})


def describe(exc: Exception) -> str:
    if isinstance(exc, ValidationError):
        return "; ".join(f"{'.'.join(map(str, e['loc']))}: {str(e['msg']).removeprefix('Value error, ')}" for e in exc.errors())
    return str(exc)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--commit", action="store_true", help="actually write to the database (default is a dry run)")
    ap.add_argument("--strict", action="store_true", help="write nothing if any row is invalid or duplicate")
    ap.add_argument("--date-format", default="%Y-%m-%d", help="strptime format for date text (default %%Y-%%m-%%d)")
    args = ap.parse_args()

    if not args.file.is_file():
        print(f"File not found: {args.file}", file=sys.stderr)
        return 2
    settings = load_settings()

    headers, rows = read_rows(args.file)
    missing = [c for c in REQUIRED if c not in headers]
    if missing:
        print(f"Missing required column(s): {', '.join(missing)}\nFound: {', '.join(h for h in headers if h)}", file=sys.stderr)
        return 2
    unknown = [h for h in headers if h and h not in REQUIRED + OPTIONAL]
    if unknown:
        print(f"Note: ignoring unrecognised column(s): {', '.join(unknown)}")

    init_db(settings.database_path)
    conn = connect(settings.database_path)
    repo = RegistryRepository(conn)
    vault = CodeVault(settings.secret_key)

    valid: list[tuple[int, RegistryWrite]] = []
    failures: list[tuple[int, str]] = []
    duplicates: list[tuple[int, str]] = []

    counts = Counter()
    parsed: list[tuple[int, RegistryWrite]] = []
    for line, row in rows:
        try:
            parsed.append((line, build_record(row, args.date_format)))
            counts[parsed[-1][1].registry_number] += 1
        except (ValidationError, ValueError) as exc:
            failures.append((line, describe(exc)))

    for line, rec in parsed:
        if counts[rec.registry_number] > 1:
            # Every copy of a duplicated number is refused: we cannot know which one is right.
            duplicates.append((line, f"{rec.registry_number} appears more than once in the file"))
        elif repo.get_by_number(rec.registry_number):
            duplicates.append((line, f"{rec.registry_number} already exists in the database"))
        else:
            valid.append((line, rec))

    imported = 0
    blocked = args.strict and (failures or duplicates)
    if args.commit and not blocked:
        for line, rec in valid:
            try:
                repo.create(to_storage(rec, vault))
                imported += 1
            except DuplicateRegistryNumber:
                duplicates.append((line, f"{rec.registry_number} already exists in the database"))
    conn.close()

    print(f"\nFile: {args.file}   Rows read: {len(rows)}")
    for label, items in (("INVALID", failures), ("DUPLICATE", duplicates)):
        for line, why in items:
            print(f"  [{label}] row {line}: {why}")
    print()
    if args.commit:
        print(f"Successfully imported: {imported}")
    else:
        print(f"DRY RUN - would import: {len(valid)}   (re-run with --commit to write)")
    print(f"Failed: {len(failures)}")
    print(f"Duplicates: {len(duplicates)}")
    if blocked:
        print("\n--strict: problems found, NOTHING was imported.")
    return 1 if (failures or duplicates) else 0


if __name__ == "__main__":
    raise SystemExit(main())
